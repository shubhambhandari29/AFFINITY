import asyncio
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from core.models.loss_run.affinity.affinity_loss_run import AffinityLossRunSelection
from services.loss_run import loss_run_job_repository as repository
from services.loss_run import loss_run_job_service as jobs
from services.loss_run import loss_run_worker as worker_module
from services.loss_run.affinity import affinity_loss_run_service as service

ROOT = Path(__file__).resolve().parents[4]
PROGRAM = {
    "programId": 3,
    "programName": "Example Program",
    "agentCodes": ["5701083", "5701212"],
}


@pytest.mark.parametrize("amounts,expected", [([19000, 21000, 18000], "Yes"), ([25000, 25000], "No"), ([30000, -10000], "No")])
def test_revised_workbook_claim_totals_and_cover(amounts, expected):
    template = (ROOT / "AffinityLossRunTemplate_Revised.xlsx").read_bytes()
    workbook = load_workbook(BytesIO(template))
    headers = [cell.value for cell in workbook["Claim Details"][1]]
    workbook.close()
    records = []
    for exposure, amount in enumerate(amounts, 1):
        row = dict.fromkeys(headers)
        row.update({"Claim Number": "claim-1", "Exposure": exposure,
                    "Incurred w/o ALAE": amount})
        records.append(row)
    program = dict(PROGRAM, agentNames=["5701083: First Agent", "5701212: Second Agent"])
    output = load_workbook(BytesIO(service.create_affinity_workbook(
        template, program, records, [], date(2026, 10, 29)
    )))
    assert output["Cover Page"]["B3"].value == "5701083: First Agent\n5701212: Second Agent"
    assert output["Cover Page"]["B4"].value == PROGRAM["programName"]
    assert output["Cover Page"]["B5"].value.date() == date(2026, 10, 29)
    assert len(output["Cover Page"]._images) == 1
    for values in output["Claim Details"].iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        assert row["Whole Claim Incurred w/o ALAE"] == sum(amounts)
        assert row["Claims Above $50k"] == expected
    output.close()
    with pytest.raises(ValueError, match="Duplicate Affinity"):
        service.create_affinity_workbook(template, program, [records[0], records[0]], [], date.today())
    records[0]["Incurred w/o ALAE"] = None
    with pytest.raises(ValueError, match="updated view"):
        service.create_affinity_workbook(template, program, records, [], date.today())


def test_programs_group_and_deduplicate_agents(monkeypatch):
    row = dict(
        ProgramId=3,
        ProgramName="Example Program",
        AcctStatus="Active",
        LossRunDistFreq="Quarterly",
    )
    query = AsyncMock(
        return_value=[dict(row, AgentCode=code) for code in ["001", "002", "001"]]
    )
    monkeypatch.setattr(service, "run_raw_query_async", query)
    result = asyncio.run(service.get_affinity_loss_run_programs())
    assert len(result) == 1
    assert result[0]["agentCodes"] == ["001", "002"]
    assert "a.LossRunReq = 'Yes'" in query.call_args.args[0]


@pytest.mark.parametrize("ids,expected", [([3, 3], ["3"]), (None, None)])
def test_create_job_uses_program_ids_and_affinity_domain(monkeypatch, ids, expected):
    monkeypatch.setattr(
        service, "get_affinity_loss_run_programs", AsyncMock(return_value=[PROGRAM])
    )
    create = AsyncMock(return_value=(uuid4(), True))
    monkeypatch.setattr(service, "create_job", create)
    result = asyncio.run(
        service.create_affinity_loss_run_job(ids, {"user": {"email": "tester"}})
    )
    assert result["status"] == "queued"
    assert create.call_args.args[2] == expected
    assert create.call_args.kwargs == {"business_type": "affinity"}


def test_unknown_program_not_queued(monkeypatch):
    monkeypatch.setattr(
        service, "get_affinity_loss_run_programs", AsyncMock(return_value=[PROGRAM])
    )
    create = AsyncMock()
    monkeypatch.setattr(service, "create_job", create)
    with pytest.raises(HTTPException) as error:
        asyncio.run(service.create_affinity_loss_run_job([99], {}))
    assert error.value.status_code == 400
    create.assert_not_awaited()


@pytest.mark.parametrize(
    "payload",
    [
        {"programIds": []},
        {"programIds": [-1]},
        {"programIds": [3], "lossDateFrom": "2004-01-01"},
    ],
)
def test_selection_rejects_invalid_or_unsupported_options(payload):
    with pytest.raises(ValidationError):
        AffinityLossRunSelection(**payload)


def test_template_headers_match_query_outputs():
    workbook = load_workbook(ROOT / "AffinityLossRunTemplate_Revised.xlsx")
    deployment = (ROOT / "services" / "sql" / "affinity_loss_run_views.sql").read_text()
    for sheet_name, view_name, runtime_query in [
        ("Claim Details", "Affinity_Loss_Run_Claims", service.CLAIMS_QUERY),
        ("Record Only", "Affinity_Loss_Run_RecordOnly", service.RECORD_ONLY_QUERY),
    ]:
        query = deployment.split(f"CREATE OR ALTER VIEW dbo.{view_name}\nAS\n")[
            1
        ].split("\nGO")[0]
        select = re.search(
            r"SELECT\s+D.PMS_CUSTMR_NBR.*?(?=GROUP BY\s+D.PMS_CUSTMR_NBR)", query, re.S
        ).group()
        aliases = set(re.findall(r"\[([^]]+)\]", select))
        aliases |= set(re.findall(r"D\.(CLMNT_\w+)\s*,", select))
        expected = {cell.value for cell in workbook[sheet_name][1]} - {
            "Lag Time", "Whole Claim Incurred w/o ALAE", "Claims Above $50k"
        }
        assert expected <= aliases, expected - aliases
        assert "?" not in query
        assert f"FROM dbo.{view_name}" in runtime_query
        assert runtime_query.count("?") == 1
        assert "WHERE [Producing Agent Code] IN" in runtime_query
        assert "OPENJSON(?)" in runtime_query
        assert "ORDER BY" in runtime_query
        assert not re.search(
            r"\b(NVL|TRUNC|ADD_MONTHS|TO_CHAR|EXTRACT)\s*\(", query, re.I
        )
        assert "CLM_LakeHouse.dbo.DW_FINCL_LOSS_FTR_SS_F" in query
    assert "DATEADD(MONTH, -60" in deployment
    assert "NOT EXISTS" in deployment
    assert ") - 6, 1, 1)" in deployment


def test_revised_financial_sql_keeps_lakehouse_and_aggregates_expense_once():
    deployment = (ROOT / "services/sql/affinity_loss_run_views.sql").read_text()
    claims = deployment.split("CREATE OR ALTER VIEW dbo.Affinity_Loss_Run_RecordOnly")[0]
    assert "CLMAA_SpecialAccounts.dbo.DW_FINCL_LOSS_FTR_SS_F_FILTERED" not in claims
    assert "FINCL.CHNG_IN_ALAE_RSRV_AMT" in claims
    assert "TXN.COST_CATGRY_DESC <> 'State Penalties'" in claims
    assert "OR TXN.COST_CATGRY_DESC IS NULL" in claims
    assert "THEN CAST(MAX(D.SAC_ALAE_EXPENSE) AS decimal(28,2))" in claims
    assert "SUM(D.SAC_ALAE_EXPENSE)" not in claims
    # Legacy aliases remain available during the staged template rollout.
    for alias in ("Current Loss Reserve", "Paid Loss Net Salvage Subro", "Total Incurred"):
        assert f"AS [{alias}]" in claims


def test_financial_operands_are_cast_before_arithmetic_and_case():
    sql = (ROOT / "services/sql/affinity_loss_run_views.sql").read_text()
    block = sql.split("-- SUM(decimal) widens", 1)[1].split("AS [Total Incurred + ALAE]", 1)[0]
    # Remove safe CASTs and comparison-only SUMs. No unbounded decimal(38,s)
    # aggregates may feed arithmetic or CASE branches in this financial block.
    remaining = re.sub(
        r"CAST\((?:SUM|MAX)\(D\.\w+\) AS decimal\(28,2\)\)", "safe_amount", block
    )
    remaining = remaining.replace("SUM(D.FTR_ALAE_AMT2) = 0", "is_zero")
    assert not re.search(r"(?:SUM|MAX)\(D\.", remaining)


def test_workbook_preserves_financial_cents():
    from decimal import Decimal

    template = (ROOT / "AffinityLossRunTemplate_Revised.xlsx").read_bytes()
    workbook = load_workbook(BytesIO(template))
    row = dict.fromkeys(cell.value for cell in workbook["Claim Details"][1])
    workbook.close()
    row.update({
        "Claim Number": "test-claim", "Exposure": "01",
        "Incurred w/o ALAE": Decimal("17518.71"),
        "ALAE Reserve": Decimal("0.00"), "ALAE Paid": Decimal("4286.27"),
        "Total Incurred + ALAE": Decimal("21804.98"),
    })
    result = load_workbook(BytesIO(service.create_affinity_workbook(
        template, PROGRAM, [row], [], date.today()
    )), data_only=True)
    sheet = result["Claim Details"]
    values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
    for name in ("Incurred w/o ALAE", "ALAE Reserve", "ALAE Paid", "Total Incurred + ALAE"):
        assert Decimal(str(values[name])) == row[name]
    result.close()


def test_workbook_fills_revised_template_and_no_claims():
    template = (ROOT / "AffinityLossRunTemplate_Revised.xlsx").read_bytes()
    workbook = load_workbook(BytesIO(template))
    record = dict.fromkeys(cell.value for cell in workbook["Claim Details"][1])
    record.update(
        {
            "Claim Number": "01-123",
            "Exposure": 1.0,
            "Loss Date": date(2020, 1, 1),
            "Date of Notice (Report Date)": datetime(2020, 1, 4, 12),
            "Incurred w/o ALAE": 123.45,
            "Description Of Loss": "=not-a-formula",
        }
    )
    result = service.create_affinity_workbook(
        template, PROGRAM, [record], [], date(2026, 10, 26)
    )
    output = load_workbook(BytesIO(result))
    assert len(output["Cover Page"]._images) == 1
    assert output["Cover Page"]["B2"].value == "5701083, 5701212"
    assert output["Cover Page"]["B4"].value == PROGRAM["programName"]
    sheet = output["Claim Details"]
    cells = {
        sheet.cell(1, i).value: sheet.cell(2, i) for i in range(1, sheet.max_column + 1)
    }
    assert cells["Lag Time"].value == 3
    assert cells["Exposure"].value == "01"
    assert cells["Incurred w/o ALAE"].value == 123.45
    assert cells["Description Of Loss"].data_type == "s"
    assert output["Record Only"]["A2"].value == "No record-only claims found"
    assert all(not column.hidden for column in sheet.column_dimensions.values())
    empty = load_workbook(
        BytesIO(
            service.create_affinity_workbook(template, PROGRAM, [], [], date.today())
        )
    )
    assert empty["Claim Details"]["A2"].value == "No claims found"


def test_generation_queries_selected_agents_and_uploads(monkeypatch, capsys):
    monkeypatch.setattr(
        service, "get_affinity_loss_run_programs", AsyncMock(return_value=[PROGRAM])
    )
    query = AsyncMock(return_value=[])
    monkeypatch.setattr(service, "run_raw_query_async", query)
    storage = MagicMock()
    storage.download_template.return_value = (
        ROOT / "AffinityLossRunTemplate_Revised.xlsx"
    ).read_bytes()
    storage.upload_report.return_value = "/Volumes/report.xlsx"
    monkeypatch.setattr(service, "DatabricksLossRunStorage", lambda: storage)
    phase, customers, result = AsyncMock(), AsyncMock(), AsyncMock()
    asyncio.run(
        service.generate_affinity_loss_runs(
            ["3"], on_phase=phase, on_customers=customers, on_result=result
        )
    )
    assert query.await_count == 2
    assert query.call_args_list[0].args == (
        service.CLAIMS_QUERY,
        ['["5701083", "5701212"]'],
    )
    assert query.call_args_list[1].args == (
        service.RECORD_ONLY_QUERY,
        ['["5701083", "5701212"]'],
    )
    storage.download_template.assert_called_once_with("affinity")
    assert re.fullmatch(
        r"3_Example Program_\d{4}_\d{2}_\d{2}\.xlsx",
        storage.upload_report.call_args.args[0],
    )
    assert re.fullmatch(r"[a-f0-9]{32}", storage.upload_report.call_args.kwargs["subdirectory"])
    result.assert_awaited_once_with("3", True, None, "/Volumes/report.xlsx")
    output = capsys.readouterr().out
    for stage in ("template", "claims_query", "record_only_query", "workbook", "upload", "program_total"):
        assert f"stage={stage} seconds=" in output
    assert "rows=0" in output
    assert "Example Program" not in output


def test_sac_job_is_not_visible_through_affinity(monkeypatch):
    monkeypatch.setattr(
        jobs, "get_job", AsyncMock(return_value={"BusinessType": "sac"})
    )
    with pytest.raises(HTTPException) as error:
        asyncio.run(jobs.get_loss_run_job(uuid4(), "affinity"))
    assert error.value.status_code == 404


def test_repository_persists_affinity_and_filters_jobs(monkeypatch):
    connection = MagicMock()
    connection.__enter__.return_value = connection
    cursor = connection.cursor.return_value
    cursor.fetchone.return_value = None
    monkeypatch.setattr(repository, "db_connection", lambda: connection)
    asyncio.run(
        repository.create_job("selected", "tester", ["3"], business_type="affinity")
    )
    assert cursor.execute.call_args.args[-1] == "affinity"
    assert "BusinessType" in cursor.execute.call_args.args[0]
    cursor.description = [("JobId",)]
    cursor.fetchall.return_value = []
    asyncio.run(repository.get_jobs("affinity"))
    assert "WHERE BusinessType = ?" in cursor.execute.call_args.args[0]
    assert cursor.execute.call_args.args[-1] == "affinity"


def test_worker_routes_affinity_without_calling_sac(monkeypatch):
    for name in [
        "complete_job",
        "fail_job",
        "get_account_numbers",
        "generate_loss_runs",
        "generate_affinity_loss_runs",
    ]:
        monkeypatch.setattr(worker_module, name, AsyncMock())
    worker_module.get_account_numbers.return_value = ["3"]
    job_id = uuid4()
    worker = worker_module.LossRunWorker()
    asyncio.run(
        worker._process_job(
            {"JobId": job_id, "JobType": "selected", "BusinessType": "affinity"}
        )
    )
    worker_module.generate_loss_runs.assert_not_awaited()
    assert worker_module.generate_affinity_loss_runs.call_args.args == (["3"],)
    worker_module.complete_job.assert_awaited_once_with(job_id, worker.worker_id)


def test_query_failure_is_not_no_claims_and_other_program_continues(monkeypatch):
    programs = [PROGRAM, dict(PROGRAM, programId=33)]
    monkeypatch.setattr(
        service, "get_affinity_loss_run_programs", AsyncMock(return_value=programs)
    )
    monkeypatch.setattr(
        service,
        "run_raw_query_async",
        AsyncMock(side_effect=[RuntimeError("DB failed"), [], []]),
    )
    storage = MagicMock()
    storage.download_template.return_value = (
        ROOT / "AffinityLossRunTemplate_Revised.xlsx"
    ).read_bytes()
    storage.upload_report.return_value = "/Volumes/second.xlsx"
    monkeypatch.setattr(service, "DatabricksLossRunStorage", lambda: storage)
    result = AsyncMock()
    asyncio.run(
        service.generate_affinity_loss_runs(
            None,
            on_phase=AsyncMock(),
            on_customers=AsyncMock(),
            on_result=result,
        )
    )
    assert result.await_count == 2
    assert result.call_args_list[0].args[:2] == ("3", False)
    assert result.call_args_list[1].args == ("33", True, None, "/Volumes/second.xlsx")
    storage.upload_report.assert_called_once()
