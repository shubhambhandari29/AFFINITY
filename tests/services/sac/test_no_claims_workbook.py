from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from services.loss_run.claim_review_workbook import create_claim_review_workbook
from services.loss_run.loss_run_service import _create_workbook
from services.loss_run.workbook_template import NO_CLAIMS_MESSAGE


@pytest.mark.parametrize("cutoff", [None, date(2004, 9, 1)])
@pytest.mark.parametrize("review", [False, True])
def test_empty_reports_clear_sample_data_and_keep_cover_and_headers(review, cutoff):
    root = Path(__file__).resolve().parents[3]
    template = root / ("SACClaimReviewTemplate.xlsx" if review else "SACLossRunTemplate.xlsx")
    builder = create_claim_review_workbook if review else _create_workbook
    content = builder([], "00123", "=Example Customer", template.read_bytes(),
                      cutoff, date(2026, 10, 14))
    with ZipFile(BytesIO(content)) as archive:
        assert not any("pivotTable" in name or "pivotCache" in name
                       or name.startswith("xl/charts/") for name in archive.namelist())
    workbook = load_workbook(BytesIO(content))
    original = load_workbook(template)
    try:
        cover = workbook["Cover Page"]
        assert cover["B2"].value == "00123"
        assert cover["B3"].value == "=Example Customer"
        assert cover["B3"].data_type == "s"
        assert cover["B4"].value == "10/14/2026"
        assert len(cover._images) == len(original["Cover Page"]._images)
        headers = ({"Review": 4, "Claims Details": 1} if review else
                   {"Claims Data": 1, "Record Only": 1,
                    "Summary By Policy Year": 6, "Chart": 1})
        assert set(workbook.sheetnames) == {"Cover Page", *headers}
        for name, row in headers.items():
            sheet = workbook[name]
            assert [c.value for c in sheet[row]] == [c.value for c in original[name][row]]
            values = [c.value for cells in sheet.iter_rows(min_row=row + 1)
                      for c in cells if c.value is not None]
            assert values == [NO_CLAIMS_MESSAGE]
            assert not sheet._pivots
            assert not sheet._charts
    finally:
        workbook.close()
        original.close()
