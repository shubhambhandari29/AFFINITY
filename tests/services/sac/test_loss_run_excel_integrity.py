"""Check serialized OOXML, not only whether openpyxl can reload its own output."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.pivot.cache import (
    CacheDefinition,
    CacheField,
    CacheSource,
    WorksheetSource,
)
from openpyxl.pivot.fields import DateTimeField, Missing, Number, Text
from openpyxl.pivot.table import (
    DataField,
    Location,
    PivotField,
    RowColField,
    RowColItem,
    TableDefinition,
)

from services.loss_run.pivot_cache import shared_items
from services.loss_run.report_cover import update_report_cover
from services.loss_run.standard_report_summary import populate_standard_summaries

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.mark.parametrize(
    "items,expected",
    [
        ([Number(v=1), Number(v=2)], (False, False, True, True, False)),
        ([Number(v=1.5)], (False, False, True, False, False)),
        ([Text(v="A")], (True, False, False, False, False)),
        ([Text(v="A"), Number(v=1)], (True, True, True, False, False)),
        ([Missing(), Number(v=1)], (True, False, True, True, False)),
        ([DateTimeField(v=datetime(2020, 1, 1))], (False, False, False, False, True)),
        ([], (False, False, False, False, False)),
    ],
)
def test_cache_type_flags_are_explicit_and_match_values(items, expected):
    node = shared_items(items).to_tree()
    names = (
        "containsSemiMixedTypes",
        "containsMixedTypes",
        "containsNumber",
        "containsInteger",
        "containsDate",
    )
    assert tuple(node.get(name) == "1" for name in names) == expected
    assert all(name in node.attrib for name in names)


def test_cache_marks_long_claim_descriptions():
    assert shared_items([Text(v="x" * 300)]).to_tree().get("longText") == "1"


def test_claim_review_uses_numeric_cache_and_no_stale_item_formats():
    from services.loss_run.claim_review_workbook import _populate_pivot

    wb = template()
    pivot = wb["Chart"]._pivots[0]
    _populate_pivot(
        wb,
        pivot,
        [
            {
                "Policy Year": 2004,
                "Claim Number": "A",
                "Distinct Claim Helper": 1,
                "Total Incurred": Decimal("12.50"),
            }
        ],
    )
    field = pivot.cache.cacheFields[3].sharedItems
    assert field.containsNumber and not field.containsSemiMixedTypes
    assert field._fields[0].v == 12.5
    assert pivot.colGrandTotals and not pivot.formats
    assert not pivot.pivotFields[3].items


def template():
    wb = Workbook()
    wb.active.title = "Cover Page"
    headers = ["Policy Year", "Claim Number", "Distinct Claim Helper", "Total Incurred"]
    cache = CacheDefinition(
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(sheet="Claims Data", ref="A1:D2"),
        ),
        cacheFields=[CacheField(name=name) for name in headers],
    )
    for name, start in [("Summary By Policy Year", 6), ("Chart", 1)]:
        sheet = wb.create_sheet(name)
        for col, heading in enumerate(
            ["Policy Year", "Number of Claims", "Total Incurred"], 1
        ):
            sheet.cell(start, col, heading)
        pivot = TableDefinition(
            name=name.replace(" ", ""),
            cacheId=1,
            dataCaption="Values",
            location=Location(
                ref=f"A{start}:C{start + 1}",
                firstHeaderRow=0,
                firstDataRow=1,
                firstDataCol=1,
            ),
            pivotFields=[
                PivotField(axis="axisRow"),
                PivotField(),
                PivotField(dataField=True),
                PivotField(dataField=True),
            ],
            rowFields=[RowColField(x=0)],
            colFields=[RowColField(x=-2)],
            colItems=[RowColItem(i=0), RowColItem(i=1)],
            dataFields=[
                DataField(name="Number of Claims", fld=2, subtotal="sum"),
                DataField(name="Total Incurred", fld=3, subtotal="sum"),
            ],
            colGrandTotals=False,
        )
        pivot.cache = cache
        sheet.add_pivot(pivot)
    return wb


@pytest.mark.parametrize(
    "claims",
    [
        [],
        [
            {"Policy Year": 2004, "Claim Number": "A", "Total Incurred": 10},
            {"Policy Year": 2004, "Claim Number": "A", "Total Incurred": 20},
            {"Policy Year": 2005, "Claim Number": "B", "Total Incurred": 40},
        ],
    ],
)
def test_summary_pivots_keep_valid_indexes_and_totals(claims):
    wb = template()
    populate_standard_summaries(wb, claims)
    output = BytesIO()
    wb.save(output)
    with ZipFile(output) as archive:
        cache = ET.fromstring(archive.read("xl/pivotCache/pivotCacheDefinition1.xml"))
        fields = cache.find(S + "cacheFields")
        records = ET.fromstring(archive.read("xl/pivotCache/pivotCacheRecords1.xml"))
        assert int(cache.get("recordCount")) == len(records) == len(claims)
        for record in records:
            assert len(record) == len(fields)
            for index, item in enumerate(record):
                assert int(item.get("v")) < len(fields[index].find(S + "sharedItems"))
        for name in archive.namelist():
            if name.startswith("xl/pivotTables/pivotTable") and name.endswith(".xml"):
                pivot = ET.fromstring(archive.read(name))
                assert pivot.get("colGrandTotals") == "1"
                assert pivot.find(S + "rowItems")[-1].get("t") == "grand"
                for field in pivot.find(S + "pivotFields"):
                    if not field.get("axis"):
                        assert field.find(S + "items") is None
    result = load_workbook(output)
    assert result["Chart"].cell(result["Chart"].max_row, 3).value == (
        70 if claims else 0
    )
    assert result["Chart"].cell(result["Chart"].max_row, 2).value == (
        2 if claims else 0
    )
    result.close()


def test_client_repair_sample_drawing_is_rebuilt_without_invalid_geometry():
    sample = (
        Path(__file__).resolve().parents[3]
        / "F W Webb Company Inc_From_2004_09_01_2026_10_01.xlsx"
    )
    if not sample.exists():
        pytest.skip("Client repair sample is not available")
    wb = load_workbook(sample)
    cover = wb["Cover Page"]
    assert len(cover._images) == 1
    old_anchor = cover._images[0].anchor
    update_report_cover(wb, date(2004, 9, 1), date(2026, 10, 1))
    assert cover._images[0].anchor._from == old_anchor._from
    assert cover._images[0].anchor.to == old_anchor.to
    output = BytesIO()
    wb.save(output)
    with ZipFile(output) as archive:
        drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
        for geometry in drawing.iter(A + "prstGeom"):
            assert all(child.tag.startswith(A) for child in geometry)
        assert len(list(drawing.iter(A + "blip"))) == 1
    result = load_workbook(output)
    assert len(result["Cover Page"]._images) == 1
    result.close()


def test_full_client_sample_keeps_all_claims_and_interactive_pivots():
    import csv

    from services.loss_run.loss_run_service import _create_workbook

    root = Path(__file__).resolve().parents[3]
    sample = root / "F W Webb Company Inc_From_2004_09_01_2026_10_01.xlsx"
    data = root / "PREPRD_9.csv"
    if not sample.exists() or not data.exists():
        pytest.skip("Client repair sample and source extract are not available")
    wb = load_workbook(sample)
    headers = [f.name for f in wb["Chart"]._pivots[0].cache.cacheFields][:-1]
    wb.close()
    with data.open(encoding="utf-8-sig", newline="") as stream:
        records = [
            dict(zip(headers, [None if v == "NULL" else v for v in row], strict=True))
            for row in csv.reader(stream)
        ]
    for record in records:
        for header in [*headers[18:29], "Total Incurred"]:
            if record[header] is not None:
                record[header] = Decimal(record[header])
        record["Policy Year"] = int(record["Policy Year"])
    output = _create_workbook(
        records,
        "0033165294",
        "F W Webb Company Inc",
        sample.read_bytes(),
        date(2004, 9, 1),
        date(2026, 10, 1),
    )
    result = load_workbook(BytesIO(output))
    actual = {
        (row[2], row[3] or None)
        for name in ("Claims Data", "Record Only")
        for row in list(result[name].values)[1:]
        if row[2]
    }
    assert actual == {(r["Claim Number"], r["Exposure"]) for r in records}
    assert len(actual) == 2479
    for name in ("Summary By Policy Year", "Chart"):
        pivot = result[name]._pivots[0]
        assert pivot.cache.recordCount == 2248
        assert pivot.colGrandTotals
        assert pivot.cache.enableRefresh
        assert not pivot.formats
    result.close()
