from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from services.loss_run.claim_review_workbook import create_claim_review_workbook
from services.loss_run.claim_threshold import (
    THRESHOLD_FIELD,
    THRESHOLD_NOT_CONFIGURED,
    claim_threshold_flags,
    parse_threshold,
)
from services.loss_run.loss_run_service import _create_workbook

ROOT = Path(__file__).resolve().parents[3]


def claim(number, exposure, reserve, paid=0, **extra):
    return {
        "Customer Number": "00123",
        "Claim Number": number,
        "Exposure": exposure,
        "Record Only Indicator": "N",
        "Policy Year": 2025,
        "Claim Status": "Open",
        "Outstanding Loss Reserve": Decimal(str(reserve)),
        "Total Paid Loss Net Salvage/Subro/Loss Recovery": Decimal(str(paid)),
        "Total Incurred": Decimal(str(reserve)) + Decimal(str(paid)),
        "Claims above 50K": 1,  # Deliberately wrong old SQL feature indicator.
        "ALAE Paid": 100000,
        **extra,
    }


def test_threshold_combines_exposures_and_includes_equality_without_expenses():
    records = [
        claim("A", 1, 6000),
        claim("A", 2, 1000, 3000),
        claim("B", 1, 9999.99),
        claim("RO", 1, 50000, **{"Record Only Indicator": "Y"}),
    ]
    assert claim_threshold_flags(records, Decimal(10000)) == {
        "A": "Meets threshold",
        "B": "Below threshold",
    }
    assert records[0]["Claims above 50K"] == 1


@pytest.mark.parametrize(
    "value,expected",
    [
        (50000, Decimal(50000)),
        ("$12,500.50", Decimal("12500.50")),
        (0, Decimal(0)),
        (None, None),
        ("  ", None),
    ],
)
def test_parse_stored_threshold(value, expected):
    assert parse_threshold(value) == expected


@pytest.mark.parametrize("value", [-1, "NaN", "Infinity", "invalid"])
def test_invalid_threshold_is_not_silently_replaced(value):
    with pytest.raises((ValueError, ArithmeticError)):
        parse_threshold(value)


@pytest.mark.parametrize("review", [False, True])
@pytest.mark.parametrize("threshold", [Decimal(10000), Decimal(0), None])
@pytest.mark.parametrize("empty", [False, True])
def test_both_workbooks_have_account_threshold_and_filter(review, threshold, empty):
    template = ROOT / (
        "SACClaimReviewTemplate.xlsx" if review else "SACLossRunTemplate.xlsx"
    )
    builder = create_claim_review_workbook if review else _create_workbook
    records = (
        []
        if empty
        else [claim("A", 1, 6000), claim("A", 2, 4000), claim("B", 1, 9999.99)]
    )
    output = builder(
        records,
        "00123",
        "Example",
        template.read_bytes(),
        None,
        date(2026, 10, 19),
        threshold,
    )
    workbook = load_workbook(BytesIO(output))
    try:
        assert workbook["Cover Page"]["B8"].value == (
            threshold if threshold is not None else THRESHOLD_NOT_CONFIGURED
        )
        sheet = workbook["Review" if review else "Summary By Policy Year"]
        caption = (
            THRESHOLD_NOT_CONFIGURED
            if threshold is None
            else f"Claims at/above ${threshold:,.0f}"
        )
        assert sheet.cell(1 if review else 4, 1).value == caption
        if empty:
            assert not sheet._pivots
            return
        pivot = sheet._pivots[0]
        headers = [field.name for field in pivot.cache.cacheFields]
        index = headers.index(THRESHOLD_FIELD)
        page = next(page for page in pivot.pageFields if page.fld == index)
        assert page.item is None  # All claims initially; the user chooses the filter.
        assert page.cap == caption
        assert pivot.pivotFields[index].axis == "axisPage"
        assert index not in [field.x for field in pivot.rowFields]
        cache_values = pivot.cache.cacheFields[index].sharedItems._fields
        flags = [cache_values[row._fields[index].v].v for row in pivot.cache.records.r]
        expected = (
            [THRESHOLD_NOT_CONFIGURED] * (2 if review else 3)
            if threshold is None
            else ["Meets threshold"] * (2 if review else 3)
            if threshold == 0
            else ["Meets threshold", "Below threshold"]
            if review
            else ["Meets threshold", "Meets threshold", "Below threshold"]
        )
        assert flags == expected
        details = workbook["Claims Details" if review else "Claims Data"]
        detail_headers = [cell.value for cell in details[1]]
        col = detail_headers.index(THRESHOLD_FIELD) + 1
        assert details.column_dimensions[get_column_letter(col)].hidden
        assert details.max_row == 4
        assert "Claims above 50K" not in detail_headers
        assert "Claim Above 25K" not in detail_headers
        if review:
            assert (
                sheet["I7"].value == 19999.99
            )  # Both claims remain in the initial total.
            assert [c.value for c in sheet[4]][-1] == "Total"
        with ZipFile(BytesIO(output)) as archive:
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for name in archive.namelist():
                if name.startswith("xl/pivotTables/pivotTable") and name.endswith(
                    ".xml"
                ):
                    root = ET.fromstring(archive.read(name))
                    cache_id = root.get("cacheId")
                    assert cache_id is not None
                    fields = root.find(ns + "pivotFields")
                    for entry in root.findall(ns + "pageFields/" + ns + "pageField"):
                        assert fields[int(entry.get("fld"))].get("axis") == "axisPage"
    finally:
        workbook.close()
