import asyncio
import sqlite3
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from core.models.loss_run.loss_run import LossRunOptions
from services.loss_run import loss_run_worker
from services.loss_run.claim_review_workbook import create_claim_review_workbook
from services.loss_run.loss_run_service import EXTENDED_HISTORY_QUERY, _create_workbook

ROOT = Path(__file__).resolve().parents[3]


def test_old_payload_cannot_silently_become_default_history():
    with pytest.raises(ValidationError, match="Use lossDateFrom"):
        LossRunOptions(policyEffectiveDateFrom="2004-10-01")
    assert LossRunOptions(lossDateFrom="2004-10-01").lossDateFrom == date(2004, 10, 1)
    assert LossRunOptions().lossDateFrom is None


def test_query_date_boundaries_use_occurrence_not_policy_start():
    # Execute the actual HAVING predicate with only SQL Server's DATEADD
    # translated for SQLite. This is a rule test, not a SQL Server integration test.
    predicate = EXTENDED_HISTORY_QUERY.rsplit("HAVING", 1)[1].replace(
        "DATEADD(DAY, 1, @ReportThroughDate)", "@NextDay"
    )
    db = sqlite3.connect(":memory:")
    db.execute(
        "CREATE TABLE claims (id TEXT, CLM_STATUS TEXT, DT_OF_LOSS TEXT, POL_EFF_DT TEXT, FTR_CHNG_IN_OSLS_AMT INT)"
    )
    db.executemany(
        "INSERT INTO claims VALUES (?, ?, ?, ?, ?)",
        [
            ("before", "closed", "2004-09-30", "2004-01-01", 0),
            ("boundary", "closed", "2004-10-01", "2004-01-01", 0),
            ("after", "Closed", "2004-12-01", "2004-01-01", 0),
            ("today", "closed", "2026-09-21 23:59:59", "2026-01-01", 0),
            ("future", "closed", "2026-09-22", "2026-01-01", 0),
            ("old-open", "open", "1990-01-01", "1990-01-01", 0),
            ("old-reserve", "closed", "1990-01-01", "1990-01-01", 5),
            ("no-loss", "closed", None, "2020-01-01", 0),
        ],
    )
    actual = {
        r[0]
        for r in db.execute(
            "SELECT id FROM claims D GROUP BY id HAVING " + predicate,
            {"LossDateFrom": "2004-10-01", "NextDay": "2026-09-22"},
        )
    }
    db.close()
    assert actual == {"boundary", "after", "today", "old-open", "old-reserve"}


def test_old_queued_policy_job_requires_resubmission(monkeypatch):
    generate = AsyncMock()
    fail = AsyncMock()
    monkeypatch.setattr(loss_run_worker, "generate_loss_runs", generate)
    monkeypatch.setattr(loss_run_worker, "fail_job", fail)
    worker = loss_run_worker.LossRunWorker()
    asyncio.run(
        worker._process_job(
            {
                "JobId": uuid4(),
                "JobType": "all",
                "PolicyEffectiveDateFrom": date(2004, 1, 1),
            }
        )
    )
    generate.assert_not_awaited()
    assert "Submit a new job" in fail.call_args.args[2]


def test_standard_summary_values_and_caches_match_claims():
    template = ROOT / "F W Webb Company Inc_From_2004_09_01_2026_09_18.xlsx"
    base = {
        "Customer Number": "TEST-CUSTOMER",
        "Record Only Indicator": "N",
        "Claim Number": "C1",
        "Policy Year": 2004,
        "Policy": "P1",
        "Policy Effective Date": "09/01/2004",
        "Loss Date": "10/01/2004",
        "Outstanding Loss Reserve": 20,
        "Total Paid Loss Net Salvage/Subro/Loss Recovery": 30,
        "ALAE Reserve": 2,
        "ALAE Paid": 3,
        "Total Incurred + ALAE": 55,
        "Total Incurred": 50,
    }
    records = [
        {**base, "Exposure": 1},
        {**base, "Exposure": 2},
        {**base, "Claim Number": "C2", "Exposure": 1, "Policy Year": 2005},
        {**base, "Claim Number": "RO", "Record Only Indicator": "Y", "Exposure": 1},
    ]
    output = _create_workbook(
        records,
        "TEST-CUSTOMER",
        "Example",
        template.read_bytes(),
        date(2004, 10, 1),
        date(2026, 9, 21),
    )
    wb = load_workbook(BytesIO(output))
    summary = wb["Summary By Policy Year"]
    assert [c.value for c in summary[summary.max_row]][9:] == [60, 90, 6, 9, 165]
    assert list(wb["Chart"].values) == [
        ("Row Labels", "Number of Claims", "Sum of Total Incurred"),
        (2004, 1, 100),
        (2005, 1, 50),
        ("Grand Total", 2, 150),
    ]
    for name in ("Summary By Policy Year", "Chart"):
        pivot = wb[name]._pivots[0]
        assert pivot.cache.recordCount == 3
        assert pivot.cache.refreshOnLoad is False
        assert pivot.cache.cacheSource.worksheetSource.sheet == "Summary Source"
    assert not any(c.value == "TEST" for s in wb for row in s for c in row)
    assert "10/1/2004 through 9/21/2026" in wb["Cover Page"]["A6"].value
    assert len(wb["Cover Page"]._images) == 1
    wb.close()


@pytest.mark.parametrize("cutoff", [None, date(2004, 10, 1)])
def test_claim_review_cover_describes_loss_history(cutoff):
    output = create_claim_review_workbook(
        [
            {
                "Claim Number": "C1",
                "Record Only Indicator": "N",
                "Outstanding Loss Reserve": 20,
                "Total Paid Loss Net Salvage/Subro/Loss Recovery": 30,
            }
        ],
        "001",
        "Example",
        (ROOT / "SACClaimReviewTemplate.xlsx").read_bytes(),
        cutoff,
        date(2026, 9, 21),
    )
    wb = load_workbook(BytesIO(output))
    wording = wb["Cover Page"]["A6"].value
    assert (
        ("10/1/2004 through 9/21/2026" in wording)
        if cutoff
        else ("six policy years" in wording)
    )
    assert wb["Review"]["I5"].value == 50
    wb.close()
