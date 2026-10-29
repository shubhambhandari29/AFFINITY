import pytest
from openpyxl import load_workbook

from scripts.update_affinity_template import (
    FINANCIAL_COLUMNS,
    TEMPLATE,
    update_template,
)


def test_revised_layout_preserves_runtime_template_and_is_repeatable(tmp_path):
    original = TEMPLATE.read_bytes()
    output = tmp_path / "revised.xlsx"
    update_template(output_path=output)
    assert TEMPLATE.read_bytes() == original
    for attempt in range(2):
        workbook = load_workbook(output)
        cover = workbook["Cover Page"]
        assert [cover.cell(row, 1).value for row in range(2, 6)] == [
            "Producer Number", "Producing Agent Name", "Program Name", "Report Run Date"
        ]
        assert all(cover.cell(row, 2).value is None for row in range(2, 6))
        assert len(cover._images) == 1
        assert "A6:B6" in map(str, cover.merged_cells.ranges)
        assert "A7:B7" in map(str, cover.merged_cells.ranges)
        claims = [cell.value for cell in workbook["Claim Details"][1]]
        assert claims[19:30] == FINANCIAL_COLUMNS
        assert claims[30] == "Litigation Status"
        assert claims[-2:] == ["Whole Claim Incurred w/o ALAE", "Claims Above $50k"]
        for name in ("Claim Details", "Record Only"):
            sheet = workbook[name]
            headers = [cell.value for cell in sheet[1]]
            assert len(headers) == len(set(headers))
            assert "Producing Agent Name" not in headers
            assert "Record Only Indicator" not in headers
            assert "Claims above 25K" not in headers
            assert all(cell.value is None or cell.data_type == "f"
                       for row in sheet.iter_rows(min_row=2) for cell in row)
            assert not any(d.hidden for d in sheet.column_dimensions.values())
        record_headers = [cell.value for cell in workbook["Record Only"][1]]
        assert not set(record_headers) & {
            "Policy Holder Claim Number", "Product", "Manufacturer",
            "Policy Region Description", "Policy Holder Retailer Desc",
        }
        workbook.close()
        if attempt == 0:
            update_template(path=output, output_path=output)


def test_rejects_populated_report(tmp_path):
    source = tmp_path / "populated.xlsx"
    workbook = load_workbook(TEMPLATE)
    workbook["Claim Details"]["A2"] = "customer data"
    workbook.save(source)
    workbook.close()
    output = tmp_path / "output.xlsx"
    with pytest.raises(ValueError, match="contains report data"):
        update_template(path=source, output_path=output)
    assert not output.exists()
