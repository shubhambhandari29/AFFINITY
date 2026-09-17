"""Build the editable rollout checklist from its Markdown source (python-docx)."""

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Loss_Run_Production_Rollout_Checklist.md"
OUTPUT = SOURCE.with_suffix(".docx")


def inline(paragraph, text):
    for part in re.split(r"(\[[^\]]+\]\(https?://[^)]+\))", text):
        match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", part)
        if not match:
            paragraph.add_run(part)
            continue
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        properties.append(color)
        run.append(properties)
        label = OxmlElement("w:t")
        label.text = match[1]
        run.append(label)
        link.append(run)
        paragraph._p.append(link)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("D94800")
    header = section.header.paragraphs[0]
    header.text = "LOSS RUN  |  PRODUCTION ROLLOUT"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.add_run("Deployment checklist • 17 September 2026    |    Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = SOURCE.read_text().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(lines[index])
                run.font.name, run.font.size = "Consolas", Pt(8)
                index += 1
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for row_index, values in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    inline(cell.paragraphs[0], value)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
                        run.bold = row_index == 0
                props = cells[0]._tc.getparent().get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    props.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        elif line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        else:
            style = "List Bullet" if line.startswith("- ") else "Normal"
            inline(doc.add_paragraph(style=style), line[2:] if line.startswith("- ") else line)
        index += 1
    doc.core_properties.title = "Loss Run — Production Rollout Checklist"
    doc.core_properties.subject = "Production prerequisites, deployment, validation and operations"
    doc.save(OUTPUT)
    # Confirm the saved package is readable and contains the expected content.
    check = Document(OUTPUT)
    assert len(check.tables) == 5
    assert any(p.text.startswith("12. Final go/no-go") for p in check.paragraphs)
    print(f"Created {OUTPUT.name}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables")


if __name__ == "__main__":
    main()
