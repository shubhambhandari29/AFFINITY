"""Affinity programs use agent-code selection, not SAC customer selection."""

import json
import logging
import re
from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from fastapi.concurrency import run_in_threadpool
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.db_helpers import run_raw_query_async
from services.loss_run.databricks_storage_service import DatabricksLossRunStorage
from services.loss_run.excel_package import save_report_workbook
from services.loss_run.loss_run_job_repository import create_job, get_job

logger = logging.getLogger(__name__)
SQL_DIRECTORY = Path(__file__).with_name("sql")
CLAIMS_QUERY = (SQL_DIRECTORY / "affinity_claims.sql").read_text()
RECORD_ONLY_QUERY = (SQL_DIRECTORY / "affinity_record_only.sql").read_text()
DATE_COLUMNS = {
    "Policy Effective Date",
    "Policy Expiration Date",
    "Loss Date",
    "Date of Notice (Report Date)",
    "Claim Closed Date",
    "System Date",
    "RO System Date",
}


async def get_affinity_loss_run_programs() -> list[dict]:
    rows = await run_raw_query_async("""
        SELECT p.AcctAffinityProgramKey AS ProgramId, p.ProgramName,
               p.AcctStatus, p.LossRunDistFreq, a.AgentCode
        FROM dbo.tblAcctAffinityProgram p
        INNER JOIN dbo.tblAffinityAgents a ON a.ProgramName = p.ProgramName
        WHERE a.LossRunReq = 'Yes'
          AND NULLIF(LTRIM(RTRIM(a.AgentCode)), '') IS NOT NULL
        ORDER BY p.ProgramName, a.AgentCode
    """)
    programs = {}
    for row in rows:
        key = int(row["ProgramId"])
        program = programs.setdefault(
            key,
            {
                "programId": key,
                "programName": row["ProgramName"],
                "accountStatus": row["AcctStatus"],
                "lossRunFrequency": row["LossRunDistFreq"],
                "agentCodes": [],
            },
        )
        code = str(row["AgentCode"]).strip()
        if code not in program["agentCodes"]:
            program["agentCodes"].append(code)
    return list(programs.values())


async def create_affinity_loss_run_job(program_ids, current_user):
    selected = None if program_ids is None else list(dict.fromkeys(program_ids))
    # Reject unknown/unconfigured selections before creating a job.
    programs = await get_affinity_loss_run_programs()
    eligible = {program["programId"] for program in programs}
    if not eligible or (selected is not None and not set(selected) <= eligible):
        raise HTTPException(
            400,
            detail={
                "error": "Select programs with at least one agent marked LossRunReq = Yes"
            },
        )
    user = current_user.get("user") or {}
    job_id, created = await create_job(
        "all" if selected is None else "selected",
        str(user.get("email") or user.get("id") or "unknown"),
        None if selected is None else [str(key) for key in selected],
        business_type="affinity",
    )
    if created:
        return {
            "jobId": job_id,
            "businessType": "affinity",
            "status": "queued",
            "message": "Affinity loss-run generation has been queued",
        }
    existing = await get_job(job_id)
    return {
        "jobId": job_id,
        "businessType": existing.get("BusinessType", "sac"),
        "status": existing["Status"],
        "message": "A loss-run job is already queued or running",
    }


def create_affinity_workbook(template_bytes, program, claims, record_only, report_date):
    workbook = load_workbook(BytesIO(template_bytes))
    try:
        cover = workbook["Cover Page"]
        cover["B2"] = ", ".join(program["agentCodes"])
        cover["B3"] = program["programName"]
        cover["B3"].data_type = "s"
        cover["B4"] = report_date
        cover["B4"].number_format = "mm/dd/yyyy"
        for name, records in (("Claim Details", claims), ("Record Only", record_only)):
            sheet = workbook[name]
            headers = [cell.value for cell in sheet[1]]
            if len(headers) != len(set(headers)) or "Lag Time" not in headers:
                raise ValueError(
                    f"{name}: upload the revised AffinityLossRunTemplate.xlsx"
                )
            styles = [copy(cell._style) for cell in sheet[2]]
            sheet.delete_rows(2, max(1, sheet.max_row - 1))
            for row_number, source in enumerate(records, 2):
                row = dict(source)
                missing = set(headers) - set(row) - {"Lag Time"}
                if missing:
                    raise ValueError(
                        f"{name}: query missing template columns: {sorted(missing)}"
                    )
                loss, notice = (
                    row.get("Loss Date"),
                    row.get("Date of Notice (Report Date)"),
                )
                if isinstance(loss, datetime):
                    loss = loss.date()
                if isinstance(notice, datetime):
                    notice = notice.date()
                row["Lag Time"] = (notice - loss).days if loss and notice else None
                exposure = row.get("Exposure")
                if exposure is not None:
                    text = str(exposure).strip()
                    row["Exposure"] = (
                        f"{int(float(text)):02d}"
                        if re.fullmatch(r"\d+(?:\.0+)?", text)
                        else text
                    )
                for column, header in enumerate(headers, 1):
                    value = row.get(header)
                    cell = sheet.cell(row_number, column, value)
                    cell._style = copy(styles[column - 1])
                    if isinstance(value, str):
                        # DB text must never become an Excel formula.
                        cell.data_type = "s"
                    if header in DATE_COLUMNS:
                        cell.number_format = "mm/dd/yyyy"
            if not records:
                sheet.cell(
                    2,
                    1,
                    "No claims found"
                    if name == "Claim Details"
                    else "No record-only claims found",
                )
            for dimension in sheet.column_dimensions.values():
                dimension.hidden = False
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}{max(2, sheet.max_row)}"
            )
        return save_report_workbook(workbook)
    finally:
        workbook.close()


async def generate_affinity_loss_runs(
    program_ids=None,
    *,
    on_phase,
    on_customers,
    on_result,
):
    await on_phase("loading_accounts")
    programs = await get_affinity_loss_run_programs()
    if program_ids is not None:
        requested = set(map(int, program_ids))
        programs = [
            program for program in programs if program["programId"] in requested
        ]
    await on_customers(
        [
            {
                "CustomerNum": str(program["programId"]),
                "CustomerName": program["programName"],
            }
            for program in programs
        ]
    )
    if program_ids is not None:
        found = {str(program["programId"]) for program in programs}
        for key in program_ids:
            if str(key) not in found:
                await on_result(
                    str(key), False, "Program no longer has eligible agents", None
                )
    if not programs:
        return
    await on_phase("loading_template")
    storage = await run_in_threadpool(DatabricksLossRunStorage)
    template = await run_in_threadpool(storage.download_template, "affinity")
    report_date = datetime.now(ZoneInfo("America/New_York")).date()
    # Unique run suffix prevents later manual runs overwriting earlier downloads.
    run_suffix = uuid4().hex
    for program in programs:
        key = str(program["programId"])
        try:
            await on_phase("querying_loss_run_data")
            params = [json.dumps(program["agentCodes"]), report_date]
            claims = await run_raw_query_async(CLAIMS_QUERY, params)
            record_only = await run_raw_query_async(RECORD_ONLY_QUERY, params)
            await on_phase("generating_reports")
            content = await run_in_threadpool(
                create_affinity_workbook,
                template,
                program,
                claims,
                record_only,
                report_date,
            )
            safe_name = re.sub(
                r'[<>:"/\\|?*\x00-\x1f]', "_", program["programName"]
            ).strip(" .")[:100]
            filename = (
                f"Affinity_{key}_{safe_name}_{report_date:%Y_%m_%d}_{run_suffix}.xlsx"
            )
            path = await run_in_threadpool(storage.upload_report, filename, content)
        except Exception:
            logger.exception("Affinity loss run failed for program %s", key)
            await on_result(
                key,
                False,
                "Report generation failed. Check application logs for details.",
                None,
            )
        else:
            await on_result(key, True, None, path)
