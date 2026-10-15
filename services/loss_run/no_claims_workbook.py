"""Build a template-based report when the query returns no matching records."""

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from services.loss_run.excel_package import save_report_workbook
from services.loss_run.report_cover import update_report_cover

NO_CLAIMS_MESSAGE = "No claims found for the selected reporting criteria"


def create_no_claims_workbook(
    customer_num: str,
    customer_name: str,
    template_bytes: bytes,
    report_type: str,
    loss_date_from: date | None,
    report_date: date | None,
) -> bytes:
    workbook = load_workbook(BytesIO(template_bytes))
    try:
        headers = (
            {"Review": 4, "Claims Details": 1}
            if report_type == "claim_review"
            else {"Claims Data": 1, "Record Only": 1,
                  "Summary By Policy Year": 6, "Chart": 1}
        )
        # Remove hidden legacy exports and cached pivots rather than retain
        # sample customers or construct empty pivot caches that Excel may repair.
        for sheet in list(workbook):
            if sheet.title not in {"Cover Page", *headers}:
                workbook.remove(sheet)
        workbook.defined_names.clear()
        for name, header_row in headers.items():
            sheet = workbook[name]
            sheet._pivots.clear()
            sheet._charts.clear()
            for row in sheet.iter_rows(min_row=header_row + 1):
                for cell in row:
                    cell.value = None
            for table in sheet.tables.values():
                left, top, right, _ = range_boundaries(table.ref)
                table.ref = (
                    f"{get_column_letter(left)}{top}:"
                    f"{get_column_letter(right)}{top + 1}"
                )
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
            sheet.auto_filter.ref = None
            sheet.cell(header_row + 2, 1, NO_CLAIMS_MESSAGE)
        report_date = report_date or datetime.now(ZoneInfo("America/New_York")).date()
        cover = workbook["Cover Page"]
        for row, value in ((2, customer_num), (3, customer_name),
                           (4, report_date.strftime("%m/%d/%Y"))):
            cover.cell(row, 2).value = value
            cover.cell(row, 2).data_type = "s"
        update_report_cover(workbook, loss_date_from, report_date)
        return save_report_workbook(workbook)
    finally:
        workbook.close()
