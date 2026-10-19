"""Account-specific claim totals and the Excel pivot threshold filter."""

from copy import copy
from decimal import Decimal

from openpyxl.pivot.table import PageField

THRESHOLD_FIELD = "Meets review threshold"
THRESHOLD_NOT_CONFIGURED = "Threshold not configured"


def parse_threshold(value) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    amount = Decimal(str(value).replace("$", "").replace(",", "").strip())
    if not amount.is_finite() or amount < 0:
        raise ValueError("Claim review threshold must be a finite, nonnegative amount")
    return amount


def claim_loss_total(record: dict) -> Decimal:
    total = sum(
        (
            Decimal(str(record.get(key) or 0))
            for key in (
                "Outstanding Loss Reserve",
                "Total Paid Loss Net Salvage/Subro/Loss Recovery",
            )
        ),
        Decimal(0),
    )
    if not total.is_finite():
        raise ValueError(f"Invalid amount for claim {record.get('Claim Number')}")
    return total


def claim_threshold_flags(
    records: list[dict], threshold: Decimal | None
) -> dict[str, str]:
    totals: dict[str, Decimal] = {}
    for record in records:
        if record.get("Record Only Indicator") == "Y":
            continue
        number = str(record.get("Claim Number") or "").strip()
        totals[number] = totals.get(number, Decimal(0)) + claim_loss_total(record)
    return {
        number: THRESHOLD_NOT_CONFIGURED
        if threshold is None
        else ("Meets threshold" if total >= threshold else "Below threshold")
        for number, total in totals.items()
    }


def configure_threshold_workbook(workbook, threshold: Decimal | None) -> None:
    """Adapt existing templates in memory, keeping field indexes and row layouts."""
    label = THRESHOLD_NOT_CONFIGURED
    if threshold is not None:
        amount = f"{threshold:,.2f}"
        if amount.endswith(".00"):
            amount = amount[:-3]
        label = f"Claims at/above ${amount}"
    cover = workbook["Cover Page"]
    cover["A8"] = "Claim review threshold"
    cover["B8"] = threshold if threshold is not None else THRESHOLD_NOT_CONFIGURED
    cover["B8"].number_format = '"$"#,##0.00'

    for sheet in workbook:
        for pivot in sheet._pivots:
            for index, field in enumerate(pivot.cache.cacheFields):
                if field.name not in {
                    "Claims above 50K",
                    "Claim Above 25K",
                    THRESHOLD_FIELD,
                }:
                    continue
                field.name = THRESHOLD_FIELD
                # The chart shares the source but has no report-filter area.
                if sheet.title not in {"Summary By Policy Year", "Review"}:
                    continue
                pivot.pivotFields[index].name = label
                pivot.pivotFields[index].axis = "axisPage"
                pages = list(pivot.pageFields or [])
                if not any(page.fld == index for page in pages):
                    pages.append(PageField(fld=index, hier=-1))
                pivot.pageFields = pages
                for page in pages:
                    if page.fld == index:
                        page.item = None
                        page.cap = label
                if sheet.title == "Review":
                    pivot.location.rowPageCount = 1
                    pivot.location.colPageCount = 1
                    row = 1
                    for target in ("A1", "B1"):
                        sheet[target]._style = copy(sheet["A4"]._style)
                else:
                    row = 4
                sheet.cell(row, 1, label)
                sheet.cell(row, 2, "(All)")
                alignment = copy(sheet.cell(row, 1).alignment)
                alignment.wrap_text = True
                sheet.cell(row, 1).alignment = alignment
                sheet.row_dimensions[row].height = max(
                    sheet.row_dimensions[row].height or 15, 30
                )
