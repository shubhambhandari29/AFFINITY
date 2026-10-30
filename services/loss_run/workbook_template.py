"""Shared template handling: cover artwork, report wording and empty reports."""

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor
from openpyxl.utils.cell import get_column_letter, range_boundaries

from services.loss_run.claim_threshold import (
    THRESHOLD_FIELD,
    configure_threshold_workbook,
)
from services.loss_run.excel_package import save_report_workbook

NO_CLAIMS_MESSAGE = "No claims found for the selected reporting criteria"


def update_report_cover(
    workbook, loss_date_from: date | None, report_date: date
) -> None:
    """Describe the actual range; policy dates remain informational columns."""
    cover = workbook["Cover Page"]
    # Rebuild drawing containers, not the image. Imported Office artwork can
    # carry shape markup that openpyxl reserializes in the wrong namespace
    # (for example avLst inside a:prstGeom), causing Excel's drawing repair.
    for image in cover._images:
        anchor = image.anchor
        if isinstance(anchor, TwoCellAnchor):
            image.anchor = TwoCellAnchor(
                editAs=anchor.editAs,
                _from=deepcopy(anchor._from),
                to=deepcopy(anchor.to),
            )
        elif isinstance(anchor, OneCellAnchor):
            image.anchor = OneCellAnchor(
                _from=deepcopy(anchor._from), ext=deepcopy(anchor.ext)
            )
    cover.cell(4, 2).value = report_date.strftime("%m/%d/%Y")
    if loss_date_from is not None:
        start = f"{loss_date_from.month}/{loss_date_from.day}/{loss_date_from.year}"
        end = f"{report_date.month}/{report_date.day}/{report_date.year}"
        cover.cell(6, 1).value = (
            f"This report contains closed claims with loss dates from {start} through {end}, "
            "inclusive, plus all open claims and claims with positive outstanding loss "
            "reserves regardless of loss date."
        )
    else:
        cover.cell(6, 1).value = (
            "This report contains information on all open claims; as well as closed "
            "claims limited to the current and prior six policy years."
        )


def create_no_claims_workbook(
    customer_num: str,
    customer_name: str,
    template_bytes: bytes,
    report_type: str,
    loss_date_from: date | None,
    report_date: date | None,
    claim_review_threshold=None,
) -> bytes:
    workbook = load_workbook(BytesIO(template_bytes))
    try:
        configure_threshold_workbook(workbook, claim_review_threshold)
        headers = (
            {"Review": 4, "Claims Details": 1}
            if report_type == "claim_review"
            else {
                "Claims Data": 1,
                "Record Only": 1,
                "Summary By Policy Year": 6,
                "Chart": 1,
            }
        )
        # Remove hidden legacy exports and cached pivots rather than retain
        # sample customers or construct empty pivot caches that Excel may repair.
        for sheet in list(workbook):
            if sheet.title not in {"Cover Page", *headers}:
                workbook.remove(sheet)
        workbook.defined_names.clear()
        for name, header_row in headers.items():
            sheet = workbook[name]
            sheet._pivots.clear()
            sheet._charts.clear()
            for cell in sheet[header_row]:
                if cell.value in {"Claims above 50K", "Claim Above 25K"}:
                    cell.value = THRESHOLD_FIELD
                    sheet.column_dimensions[get_column_letter(cell.column)].hidden = True
            for row in sheet.iter_rows(min_row=header_row + 1):
                for cell in row:
                    cell.value = None
            for table in sheet.tables.values():
                left, top, right, _ = range_boundaries(table.ref)
                table.ref = (
                    f"{get_column_letter(left)}{top}:"
                    f"{get_column_letter(right)}{top + 1}"
                )
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
            sheet.auto_filter.ref = None
            sheet.cell(header_row + 2, 1, NO_CLAIMS_MESSAGE)
        report_date = report_date or datetime.now(ZoneInfo("America/New_York")).date()
        cover = workbook["Cover Page"]
        for row, value in (
            (2, customer_num),
            (3, customer_name),
            (4, report_date.strftime("%m/%d/%Y")),
        ):
            cover.cell(row, 2).value = value
            cover.cell(row, 2).data_type = "s"
        update_report_cover(workbook, loss_date_from, report_date)
        return save_report_workbook(workbook)
    finally:
        workbook.close()
