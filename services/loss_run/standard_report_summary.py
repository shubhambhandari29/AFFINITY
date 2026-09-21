"""Populate the standard template's summary sheets and Excel pivot caches."""

from copy import copy
from datetime import date, datetime
from decimal import Decimal

from openpyxl.chart import BarChart, Reference
from openpyxl.pivot.cache import CacheField, SharedItems, WorksheetSource
from openpyxl.pivot.fields import DateTimeField, Index, Missing, Number, Text
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import FieldItem, RowColItem
from openpyxl.utils import get_column_letter


def populate_standard_summaries(workbook, claims: list[dict]) -> None:
    """Render real values immediately, retaining pivots for Excel filtering.

    Use the template's cache field order, including legacy helper fields, so
    removing Record Only Indicator from Claims Data cannot misalign the pivot.
    """
    targets = [
        workbook[name]
        for name in ("Summary By Policy Year", "Chart")
        if name in workbook.sheetnames and workbook[name]._pivots
    ]
    if not targets:
        return
    headers = [field.name for field in targets[0]._pivots[0].cache.cacheFields]
    if "Summary Source" in workbook.sheetnames:
        del workbook["Summary Source"]
    source = workbook.create_sheet("Summary Source")
    source.sheet_state = "hidden"
    source.append(headers)
    seen = set()
    values = []
    for claim in claims:
        claim = dict(claim)
        number = claim["Claim Number"]
        claim["Distinct Claim Helper"] = int(number not in seen)
        seen.add(number)
        row = [claim.get(header) for header in headers]
        values.append(row)
        source.append(row)
        for cell in source[source.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    fields = []
    indexes = [[] for _ in values]
    for column, header in enumerate(headers):
        unique = list(dict.fromkeys(row[column] for row in values))
        lookup = {value: index for index, value in enumerate(unique)}
        items = []
        for value in unique:
            if value is None:
                items.append(Missing())
            elif isinstance(value, int | float | Decimal):
                items.append(Number(v=float(value)))
            elif isinstance(value, date | datetime):
                items.append(DateTimeField(v=value))
            else:
                items.append(Text(v=str(value)))
        fields.append(
            CacheField(
                name=header,
                sharedItems=SharedItems(
                    _fields=items,
                    containsBlank=None in unique,
                    containsString=any(isinstance(v, str) for v in unique),
                    containsNumber=any(
                        isinstance(v, int | float | Decimal) for v in unique
                    ),
                    containsDate=any(isinstance(v, date | datetime) for v in unique),
                    containsNonDate=any(
                        not isinstance(v, date | datetime) for v in unique
                    ),
                ),
            )
        )
        for index, row in enumerate(values):
            indexes[index].append(lookup[row[column]])

    for sheet in targets:
        pivot = sheet._pivots[0]
        if [f.name for f in pivot.cache.cacheFields] != headers:
            raise ValueError(
                "Standard template pivots must use the same Claims Data fields"
            )
        cache = pivot.cache
        cache.cacheFields = fields
        cache.records = RecordList(
            r=[Record(_fields=[Index(v=v) for v in row]) for row in indexes]
        )
        cache.recordCount = len(values)
        cache.cacheSource.worksheetSource = WorksheetSource(
            sheet=source.title,
            ref=f"A1:{get_column_letter(len(headers))}{len(values) + 1}",
        )
        cache.refreshOnLoad = False
        cache.enableRefresh = True
        cache.missingItemsLimit = 0
        for index, field in enumerate(pivot.pivotFields):
            field.items = [
                FieldItem(x=i) for i in range(len(fields[index].sharedItems._fields))
            ]
            for name in (
                "defaultSubtotal",
                "sumSubtotal",
                "countASubtotal",
                "avgSubtotal",
                "maxSubtotal",
                "minSubtotal",
                "productSubtotal",
                "countSubtotal",
                "stdDevSubtotal",
                "stdDevPSubtotal",
                "varSubtotal",
                "varPSubtotal",
            ):
                setattr(field, name, False)
        for page in pivot.pageFields:
            page.item = None
        groups = {}
        for row, item_indexes in zip(values, indexes, strict=True):
            key = tuple(row[field.x] for field in pivot.rowFields)
            if key not in groups:
                groups[key] = (
                    [Decimal(0) for _ in pivot.dataFields],
                    [item_indexes[field.x] for field in pivot.rowFields],
                )
            totals, _ = groups[key]
            for index, field in enumerate(pivot.dataFields):
                if field.subtotal != "sum":
                    raise ValueError("Unexpected standard report pivot aggregation")
                totals[index] += Decimal(str(row[field.fld] or 0))

        start = 7 if sheet.title == "Summary By Policy Year" else 2
        styles = [copy(cell._style) for cell in sheet[start]]
        for row in sheet.iter_rows(min_row=start):
            for cell in row:
                cell.value = None
        rows = []
        row_items = []
        grand = [Decimal(0) for _ in pivot.dataFields]
        for key in sorted(groups, key=lambda k: tuple(str(v or "") for v in k)):
            totals, refs = groups[key]
            rows.append([*key, *totals])
            row_items.append(RowColItem(x=[Index(v=v) for v in refs]))
            grand = [a + b for a, b in zip(grand, totals, strict=True)]
        rows.append(["Grand Total", *([None] * (len(pivot.rowFields) - 1)), *grand])
        for index, row in enumerate(rows, start):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(index, column)
                if column <= len(styles):
                    cell._style = copy(styles[column - 1])
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"
        end = start + len(rows) - 1
        if sheet.max_row > end:
            sheet.delete_rows(end + 1, sheet.max_row - end)
        pivot.rowItems = row_items + [RowColItem(t="grand", x=[Index(v=0)])]
        pivot.location.ref = f"A{start - 1}:{get_column_letter(len(rows[-1]))}{end}"
        if sheet.title == "Chart":
            # Replace the template chart's stale cached TEST series.
            sheet._charts.clear()
            if groups:
                chart = BarChart()
                chart.title = "Total Incurred by Policy Year"
                chart.y_axis.title = "Total Incurred"
                chart.x_axis.title = "Policy Year"
                chart.add_data(
                    Reference(sheet, min_col=3, min_row=1, max_row=end - 1),
                    titles_from_data=True,
                )
                chart.set_categories(
                    Reference(sheet, min_col=1, min_row=2, max_row=end - 1)
                )
                sheet.add_chart(chart, "E2")
