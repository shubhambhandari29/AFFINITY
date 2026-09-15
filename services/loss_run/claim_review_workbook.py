from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.pivot.cache import CacheField, SharedItems, WorksheetSource
from openpyxl.pivot.fields import Index, Missing, Number, Text
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import FieldItem, RowColItem
from openpyxl.utils import get_column_letter

REVIEW_FIELDS = {
    "Claimant Name-Company": "Claimant Name",
    "Adjuster": "Adjuster",
    "Policy Symbol": "Policy Symbol",
    "Cause of Loss": "Cause Of Loss",
    "Reported Date": "Date of Notice",
    "Claim Status": "Claim Status",
    "Claim Closed Date": "Claim Closed Date",
}

# The client template uses legacy export headings; the view uses newer names.
DETAIL_FIELDS = {
    **REVIEW_FIELDS,
    "Customer Account Number": "Customer Number",
    "Feature Number": "Exposure",
    "Current Loss Reserve": "Outstanding Loss Reserve",
    "Paid Loss Net Subro/Salvage": "Total Paid Loss Net Salvage/Subro/Loss Recovery",
    "Policy Region Desc": "Policy Region Description",
    "Accident DOW": "Accident Day Of Week",
    "Total Loss Ind": "Total Loss Indicator",
    "Vehicle ID": "Vehicle ID Number",
    "RO_Ind": "Record Only Indicator",
    "Agent code": "Producing Agent Code",
    "Agent Name": "Producing Agent Name",
    "Adjuster Phone": "Adjuster Phone Number",
    "Record Source": "Record Source Code",
    "Expo_Ind": "Exposure Indicator",
}
DATE_FIELDS = {
    "Policy Effective Date",
    "Policy Expiration Date",
    "Loss Date",
    "Reported Date",
    "Claim Closed Date",
}


def _fill_template_rows(sheet, rows: list[list], start_row: int) -> None:
    """Replace sample values while extending the template's own row formatting."""
    styles = [copy(cell._style) for cell in sheet[start_row]]
    height = sheet.row_dimensions[start_row].height
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.value = None
    last_row = start_row + len(rows) - 1
    if sheet.max_row > last_row:
        sheet.delete_rows(last_row + 1, sheet.max_row - last_row)
    for index, values in enumerate(rows, start_row):
        sheet.row_dimensions[index].height = height
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column)
            cell._style = copy(styles[column - 1])
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"


def _populate_pivot(workbook, pivot, summaries: list[dict]) -> None:
    """Keep the client's pivot, replacing its sample cache and source completely.

    A hidden claim-level source ensures Excel Refresh still gives one row per
    claim, even where its exposures have different claimants or adjusters.
    """
    headers = [field.name for field in pivot.cache.cacheFields]
    source = workbook.create_sheet("Review Source")
    source.sheet_state = "hidden"
    source.append(headers)
    values = [[summary.get(header) for header in headers] for summary in summaries]
    for row in values:
        source.append(row)
    for row in source.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    indexes = [[] for _ in values]
    fields = []
    for column, header in enumerate(headers):
        unique = list(dict.fromkeys(row[column] for row in values))
        lookup = {value: index for index, value in enumerate(unique)}
        items = [
            Missing()
            if value is None
            else Number(v=float(value))
            if isinstance(value, Decimal)
            else Text(v=str(value))
            for value in unique
        ]
        fields.append(
            CacheField(
                name=header,
                sharedItems=SharedItems(
                    _fields=items,
                    containsBlank=None in unique,
                    containsString=any(isinstance(value, str) for value in unique),
                    containsNumber=any(isinstance(value, Decimal) for value in unique),
                    containsNonDate=True,
                ),
            )
        )
        for row_index, row in enumerate(values):
            indexes[row_index].append(lookup[row[column]])
        field = pivot.pivotFields[column]
        field.items = [FieldItem(x=index) for index in range(len(unique))]
        field.defaultSubtotal = False

    pivot.cache.cacheFields = fields
    pivot.cache.records = RecordList(
        r=[Record(_fields=[Index(v=index) for index in row]) for row in indexes]
    )
    pivot.cache.recordCount = len(values)
    pivot.cache.cacheSource.worksheetSource = WorksheetSource(
        sheet=source.title, ref=f"A1:{get_column_letter(len(headers))}{len(values) + 1}"
    )
    # Values and cache are already current. Avoid an unnecessary automatic
    # refresh reformatting the client layout when the file is first opened.
    pivot.cache.refreshOnLoad = False
    pivot.cache.enableRefresh = True
    pivot.cache.missingItemsLimit = 0
    pivot.rowItems = [
        RowColItem(x=[Index(v=row[field.x]) for field in pivot.rowFields])
        for row in indexes
    ] + [RowColItem(t="grand", x=[Index(v=0)])]
    pivot.location.ref = f"A3:I{len(values) + 5}"
    pivot.location.firstHeaderRow = 1


def create_claim_review_workbook(
    records: list[dict], customer_num: str, customer_name: str, template_bytes: bytes
) -> bytes:
    """Populate the converted client Claim Review template, not a new workbook."""
    workbook = load_workbook(BytesIO(template_bytes))
    try:
        review = workbook["Review"]
        details = workbook["Claims Details"]
        expected = ["Claim Number", *REVIEW_FIELDS, "Total"]
        if [cell.value for cell in review[4]] != expected or len(review._pivots) != 1:
            raise ValueError(
                "Claim Review template must contain the approved Review layout and pivot"
            )
        pivot = review._pivots[0]
        cache_headers = [field.name for field in pivot.cache.cacheFields]
        if [cache_headers[field.x] for field in pivot.rowFields] != expected[
            :-1
        ] or cache_headers[pivot.dataFields[0].fld] != "Incurred":
            raise ValueError(
                "Claim Review template has an unexpected pivot configuration"
            )

        # Hidden legacy exports contain other accounts and stale pivots, and are
        # not part of the client-facing report. Retain the actual visible sheets.
        for sheet in list(workbook):
            if sheet.title not in {
                "Cover Page",
                "Review",
                "Claims Details",
                "Accounts",
            }:
                workbook.remove(sheet)
        workbook.defined_names.clear()
        claims = {}
        detail_rows = []
        headers = [cell.value for cell in details[1]]
        seen = set()
        for record in records:
            if record.get("Record Only Indicator") == "Y":
                continue
            number = str(record.get("Claim Number") or "").strip()
            if not number:
                raise ValueError("Claim number is required for claim-level review")
            total = sum(
                (
                    Decimal(str(record[key] or 0))
                    for key in (
                        "Outstanding Loss Reserve",
                        "Total Paid Loss Net Salvage/Subro/Loss Recovery",
                    )
                ),
                Decimal(0),
            )
            if not total.is_finite():
                raise ValueError(f"Invalid amount for claim {number}")
            claim = claims.setdefault(
                number,
                {"values": {field: [] for field in REVIEW_FIELDS}, "total": Decimal(0)},
            )
            claim["total"] += total
            for field, source in REVIEW_FIELDS.items():
                value = record.get(source)
                if isinstance(value, date | datetime):
                    value = value.strftime("%m/%d/%Y")
                text = str(value).strip() if value is not None else ""
                if text and text not in claim["values"][field]:
                    claim["values"][field].append(text)

            row = []
            for header in headers:
                base = header.removesuffix("_Source")
                value = record.get(DETAIL_FIELDS.get(base, base))
                if header == "Incurred":
                    value = total
                elif header == "Claim Distinct":
                    value = int(number not in seen)
                elif header == "Claim Above 25K":
                    value = None  # Legacy helper, explicitly not requested.
                elif header == "Feature Number" and value not in (None, ""):
                    value = f"{int(value):02d}"
                elif base in DATE_FIELDS and isinstance(value, str) and value.strip():
                    value = datetime.strptime(value, "%m/%d/%Y")
                row.append(value)
            detail_rows.append(row)
            seen.add(number)
        _fill_template_rows(details, detail_rows, 2)
        details.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(detail_rows) + 1}"
        )

        summaries = [
            {
                "Claim Number": number,
                **{field: "; ".join(claim["values"][field]) for field in REVIEW_FIELDS},
                "Incurred": claim["total"],
            }
            for number, claim in sorted(claims.items())
        ]
        rows = [
            [summary[header] for header in expected[:-1]] + [summary["Incurred"]]
            for summary in summaries
        ]
        total_style = [copy(cell._style) for cell in review[7]]
        rows.append(
            [
                "Grand Total",
                *([None] * 7),
                sum((summary["Incurred"] for summary in summaries), Decimal(0)),
            ]
        )
        _fill_template_rows(review, rows, 5)
        for cell, style in zip(review[len(rows) + 4], total_style, strict=True):
            cell._style = style
        _populate_pivot(workbook, pivot, summaries)

        today = datetime.now().strftime("%m/%d/%Y")
        accounts = workbook["Accounts"]
        _fill_template_rows(accounts, [[customer_num, today, customer_name]], 2)
        cover = workbook["Cover Page"]
        for row, value in ((2, customer_num), (3, customer_name), (4, today)):
            cover.cell(row, 2).value = value
            cover.cell(row, 2).data_type = "s"
        workbook.calculation.fullCalcOnLoad = True
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
