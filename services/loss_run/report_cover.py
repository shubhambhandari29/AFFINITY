from copy import deepcopy
from datetime import date

from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor


def update_report_cover(
    workbook, loss_date_from: date | None, report_date: date
) -> None:
    """Describe the actual range; policy dates remain informational columns."""
    cover = workbook["Cover Page"]
    # Rebuild drawing containers, not the image. Imported Office artwork can
    # carry shape markup that openpyxl reserializes in the wrong namespace
    # (for example avLst inside a:prstGeom), causing Excel's drawing repair.
    for image in cover._images:
        anchor = image.anchor
        if isinstance(anchor, TwoCellAnchor):
            image.anchor = TwoCellAnchor(
                editAs=anchor.editAs,
                _from=deepcopy(anchor._from),
                to=deepcopy(anchor.to),
            )
        elif isinstance(anchor, OneCellAnchor):
            image.anchor = OneCellAnchor(
                _from=deepcopy(anchor._from), ext=deepcopy(anchor.ext)
            )
    cover.cell(4, 2).value = report_date.strftime("%m/%d/%Y")
    if loss_date_from is not None:
        start = f"{loss_date_from.month}/{loss_date_from.day}/{loss_date_from.year}"
        end = f"{report_date.month}/{report_date.day}/{report_date.year}"
        cover.cell(6, 1).value = (
            f"This report contains closed claims with loss dates from {start} through {end}, "
            "inclusive, plus all open claims and claims with positive outstanding loss "
            "reserves regardless of loss date."
        )
    else:
        cover.cell(6, 1).value = (
            "This report contains information on all open claims; as well as closed "
            "claims limited to the current and prior six policy years."
        )
