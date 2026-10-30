"""Affinity programs use agent-code selection, not SAC customer selection."""

import json
import logging
import re
from copy import copy
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from time import perf_counter
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from fastapi.concurrency import run_in_threadpool
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.db_helpers import run_raw_query_async
from services.loss_run.databricks_storage_service import DatabricksLossRunStorage
from services.loss_run.excel_package import save_report_workbook
from services.loss_run.loss_run_cancellation import LossRunCancellation, LossRunCancelled
from services.loss_run.loss_run_job_repository import create_job, get_job

logger = logging.getLogger(__name__)
CLAIMS_QUERY = """
    SELECT *
    FROM dbo.Affinity_Loss_Run_Claims
    WHERE [Producing Agent Code] IN (
        SELECT CAST([value] AS varchar(50)) COLLATE SQL_Latin1_General_CP1_CI_AS
        FROM OPENJSON(?)
    )
    ORDER BY [Claim Number], [Exposure]
"""
RECORD_ONLY_QUERY = """
    SELECT *
    FROM dbo.Affinity_Loss_Run_RecordOnly
    WHERE [Producing Agent Code] IN (
        SELECT CAST([value] AS varchar(50)) COLLATE SQL_Latin1_General_CP1_CI_AS
        FROM OPENJSON(?)
    )
    ORDER BY [Customer Number], [Claim Number], [Exposure]
"""
DATE_COLUMNS = {
    "Policy Effective Date",
    "Policy Expiration Date",
    "Loss Date",
    "Date of Notice (Report Date)",
    "Claim Closed Date",
    "System Date",
    "RO System Date",
}


async def get_affinity_loss_run_programs(
    cancellation: LossRunCancellation | None = None,
) -> list[dict]:
    query = cancellation.run_query if cancellation else run_raw_query_async
    rows = await query("""
        SELECT p.AcctAffinityProgramKey AS ProgramId, p.ProgramName,
               p.AcctStatus, p.LossRunDistFreq, a.AgentCode, a.AgentName
        FROM dbo.tblAcctAffinityProgram p
        INNER JOIN dbo.tblAffinityAgents a ON a.ProgramName = p.ProgramName
        WHERE a.LossRunReq = 'Yes'
          AND NULLIF(LTRIM(RTRIM(a.AgentCode)), '') IS NOT NULL
        ORDER BY p.ProgramName, a.AgentCode
    """)
    programs = {}
    for row in rows:
        key = int(row["ProgramId"])
        program = programs.setdefault(
            key,
            {
                "programId": key,
                "programName": row["ProgramName"],
                "accountStatus": row["AcctStatus"],
                "lossRunFrequency": row["LossRunDistFreq"],
                "agentCodes": [],
                "agentNames": [],
            },
        )
        code = str(row["AgentCode"]).strip()
        if code not in program["agentCodes"]:
            program["agentCodes"].append(code)
        name = str(row.get("AgentName") or "").strip()
        if name and f"{code}: {name}" not in program["agentNames"]:
            program["agentNames"].append(f"{code}: {name}")
    return list(programs.values())


async def create_affinity_loss_run_job(program_ids, current_user):
    selected = None if program_ids is None else list(dict.fromkeys(program_ids))
    # Reject unknown/unconfigured selections before creating a job.
    programs = await get_affinity_loss_run_programs()
    eligible = {program["programId"] for program in programs}
    if not eligible or (selected is not None and not set(selected) <= eligible):
        raise HTTPException(
            400,
            detail={
                "error": "Select programs with at least one agent marked LossRunReq = Yes"
            },
        )
    user = current_user.get("user") or {}
    job_id, created = await create_job(
        "all" if selected is None else "selected",
        str(user.get("email") or user.get("id") or "unknown"),
        None if selected is None else [str(key) for key in selected],
        business_type="affinity",
    )
    if created:
        return {
            "jobId": job_id,
            "businessType": "affinity",
            "status": "queued",
            "message": "Affinity loss-run generation has been queued",
        }
    existing = await get_job(job_id)
    return {
        "jobId": job_id,
        "businessType": existing.get("BusinessType", "sac"),
        "status": existing["Status"],
        "message": "A loss-run job is already queued or running",
    }


def create_affinity_workbook(template_bytes, program, claims, record_only, report_date):
    workbook = load_workbook(BytesIO(template_bytes))
    try:
        cover = workbook["Cover Page"]
        cover["B2"] = ", ".join(program["agentCodes"])
        cover["B2"].data_type = "s"
        revised = cover["A3"].value == "Producing Agent Name"
        if revised:
            cover["B3"] = "\n".join(program.get("agentNames", []))
            cover["B3"].data_type = "s"
            alignment = copy(cover["B3"].alignment)
            alignment.wrap_text = True
            cover["B3"].alignment = alignment
            cover.row_dimensions[3].height = 15 * max(1, len(program.get("agentNames", [])))
        program_cell, date_cell = ("B4", "B5") if revised else ("B3", "B4")
        cover[program_cell] = program["programName"]
        cover[program_cell].data_type = "s"
        cover[date_cell] = report_date
        cover[date_cell].number_format = "mm/dd/yyyy"
        claim_totals = {}
        if revised:
            # Require the aligned financial field; never substitute the old
            # Affinity Total Incurred calculation or silently fill missing values.
            exposures = {}
            for row in claims:
                claim = str(row.get("Claim Number") or "").strip()
                exposure = str(row.get("Exposure") or "").strip()
                if re.fullmatch(r"\d+(?:\.0+)?", exposure):
                    exposure = str(int(Decimal(exposure)))
                amount = row.get("Incurred w/o ALAE")
                if not claim or not exposure or amount is None:
                    raise ValueError("Revised Affinity report requires claim, exposure and Incurred w/o ALAE from the updated view")
                key = (claim, exposure)
                amount = Decimal(str(amount))
                if key in exposures:
                    # A repeated exposure can indicate join fan-out. Do not
                    # silently double-count or guess which financial row is valid.
                    raise ValueError(f"Duplicate Affinity claim/exposure: {claim}/{exposure}; verify view aggregation")
                exposures[key] = amount
                claim_totals[claim] = claim_totals.get(claim, Decimal("0")) + amount
        for name, records in (("Claim Details", claims), ("Record Only", record_only)):
            sheet = workbook[name]
            headers = [cell.value for cell in sheet[1]]
            if len(headers) != len(set(headers)) or "Lag Time" not in headers:
                raise ValueError(
                    f"{name}: upload the revised AffinityLossRunTemplate.xlsx"
                )
            styles = [copy(cell._style) for cell in sheet[2]]
            sheet.delete_rows(2, max(1, sheet.max_row - 1))
            for row_number, source in enumerate(records, 2):
                row = dict(source)
                if revised and name == "Claim Details":
                    total = claim_totals[str(row["Claim Number"]).strip()]
                    row["Whole Claim Incurred w/o ALAE"] = total
                    row["Claims Above $50k"] = "Yes" if total > Decimal("50000") else "No"
                missing = set(headers) - set(row) - {"Lag Time"}
                if missing:
                    raise ValueError(
                        f"{name}: query missing template columns: {sorted(missing)}"
                    )
                loss, notice = (
                    row.get("Loss Date"),
                    row.get("Date of Notice (Report Date)"),
                )
                if isinstance(loss, datetime):
                    loss = loss.date()
                if isinstance(notice, datetime):
                    notice = notice.date()
                row["Lag Time"] = (notice - loss).days if loss and notice else None
                exposure = row.get("Exposure")
                if exposure is not None:
                    text = str(exposure).strip()
                    row["Exposure"] = (
                        f"{int(float(text)):02d}"
                        if re.fullmatch(r"\d+(?:\.0+)?", text)
                        else text
                    )
                for column, header in enumerate(headers, 1):
                    value = row.get(header)
                    cell = sheet.cell(row_number, column, value)
                    cell._style = copy(styles[column - 1])
                    if isinstance(value, str):
                        # DB text must never become an Excel formula.
                        cell.data_type = "s"
                    if header in DATE_COLUMNS:
                        cell.number_format = "mm/dd/yyyy"
            if not records:
                sheet.cell(
                    2,
                    1,
                    "No claims found"
                    if name == "Claim Details"
                    else "No record-only claims found",
                )
            for dimension in sheet.column_dimensions.values():
                dimension.hidden = False
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}{max(2, sheet.max_row)}"
            )
        return save_report_workbook(workbook)
    finally:
        workbook.close()


async def generate_affinity_loss_runs(
    program_ids=None,
    *,
    on_phase,
    on_customers,
    on_result,
    cancellation: LossRunCancellation | None = None,
):
    query = cancellation.run_query if cancellation else run_raw_query_async
    if cancellation:
        cancellation.check()
    await on_phase("loading_accounts")
    programs = (await get_affinity_loss_run_programs(cancellation=cancellation)
                if cancellation else await get_affinity_loss_run_programs())
    if program_ids is not None:
        requested = set(map(int, program_ids))
        programs = [
            program for program in programs if program["programId"] in requested
        ]
    await on_customers(
        [
            {
                "CustomerNum": str(program["programId"]),
                "CustomerName": program["programName"],
            }
            for program in programs
        ]
    )
    if program_ids is not None:
        found = {str(program["programId"]) for program in programs}
        for key in program_ids:
            if str(key) not in found:
                await on_result(
                    str(key), False, "Program no longer has eligible agents", None
                )
    if not programs:
        return
    if cancellation:
        cancellation.check()
    await on_phase("loading_template")
    template_started = perf_counter()
    storage = await run_in_threadpool(DatabricksLossRunStorage)
    template = await run_in_threadpool(storage.download_template, "affinity")
    if cancellation:
        cancellation.check()
    # Temporary diagnostic output: visible in standalone-worker terminals too,
    # without changing application-wide logging levels or environment settings.
    print(f"[AFFINITY TIMING] stage=template seconds={perf_counter() - template_started:.2f}", flush=True)
    report_date = datetime.now(ZoneInfo("America/New_York")).date()
    # Keep run uniqueness in storage, not in the user-facing Excel filename.
    run_directory = uuid4().hex
    for program in programs:
        key = str(program["programId"])
        program_started = perf_counter()
        stage = "claims_query"
        stage_started = program_started
        try:
            if cancellation:
                cancellation.check()
            await on_phase("querying_loss_run_data")
            params = [json.dumps(program["agentCodes"])]
            stage_started = perf_counter()
            print(f"[AFFINITY TIMING] program={key} stage={stage} started", flush=True)
            claims = await query(CLAIMS_QUERY, params)
            if cancellation:
                cancellation.check()
            print(f"[AFFINITY TIMING] program={key} stage={stage} seconds={perf_counter() - stage_started:.2f} rows={len(claims)}", flush=True)
            stage = "record_only_query"
            stage_started = perf_counter()
            print(f"[AFFINITY TIMING] program={key} stage={stage} started", flush=True)
            record_only = await query(RECORD_ONLY_QUERY, params)
            if cancellation:
                cancellation.check()
            print(f"[AFFINITY TIMING] program={key} stage={stage} seconds={perf_counter() - stage_started:.2f} rows={len(record_only)}", flush=True)
            await on_phase("generating_reports")
            stage = "workbook"
            stage_started = perf_counter()
            content = await run_in_threadpool(
                create_affinity_workbook,
                template,
                program,
                claims,
                record_only,
                report_date,
            )
            if cancellation:
                cancellation.check()
            print(f"[AFFINITY TIMING] program={key} stage={stage} seconds={perf_counter() - stage_started:.2f}", flush=True)
            safe_name = re.sub(
                r'[<>:"/\\|?*\x00-\x1f]', "_", program["programName"]
            ).strip(" .")[:100]
            filename = f"{key}_{safe_name}_{report_date:%Y_%m_%d}.xlsx"
            stage = "upload"
            stage_started = perf_counter()
            path = await run_in_threadpool(
                storage.upload_report, filename, content, subdirectory=run_directory
            )
            print(f"[AFFINITY TIMING] program={key} stage={stage} seconds={perf_counter() - stage_started:.2f}", flush=True)
        except LossRunCancelled:
            raise
        except Exception:
            if cancellation:
                cancellation.check()
            print(f"[AFFINITY TIMING] program={key} stage={stage} failed_after={perf_counter() - stage_started:.2f}s", flush=True)
            logger.exception("Affinity loss run failed for program %s", key)
            await on_result(
                key,
                False,
                "Report generation failed. Check application logs for details.",
                None,
            )
        else:
            await on_result(key, True, None, path)
            # Preserve an uploaded file before acknowledging cancellation.
            if cancellation:
                cancellation.check()
        finally:
            print(f"[AFFINITY TIMING] program={key} stage=program_total seconds={perf_counter() - program_started:.2f}", flush=True)
